import io
import os
import pypdf
import docx
from pptx import Presentation

class DocumentParserTool:
    """Parses PDF, Word (docx), PowerPoint (pptx), and text files."""
    
    @staticmethod
    def extract_text(file_bytes: bytes, filename: str) -> str:
        ext = os.path.splitext(filename)[1].lower()
        
        try:
            if ext == ".pdf":
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                text_list = []
                for i, page in enumerate(reader.pages[:20]): # Up to 20 pages
                    t = page.extract_text()
                    if t:
                        text_list.append(f"--- [পৃষ্ঠা {i+1}] ---\n{t}")
                return "\n\n".join(text_list) if text_list else "[PDF ফাইলটিতে কোনো টেক্সট পাওয়া যায়নি]"
                
            elif ext == ".docx":
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                return "\n".join(paragraphs) if paragraphs else "[Word ডকুমেন্টে কোনো টেক্সট পাওয়া যায়নি]"
                
            elif ext == ".doc":
                return "[.doc ফরম্যাট সাপোর্টেড নয়। অনুগ্রহ করে ফাইলটি .docx ফরম্যাটে সেভ করে আবার পাঠান।]"
                
            elif ext in [".pptx", ".ppt"]:
                prs = Presentation(io.BytesIO(file_bytes))
                slides_text = []
                for i, slide in enumerate(prs.slides):
                    slide_lines = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_lines.append(shape.text.strip())
                    if slide_lines:
                        slides_text.append(f"--- [স্লাইড {i+1}] ---\n" + "\n".join(slide_lines))
                return "\n\n".join(slides_text) if slides_text else "[PowerPoint প্রেজেন্টেশনে কোনো টেক্সট পাওয়া যায়নি]"
                
            else:
                # Text, markdown, source code files
                return file_bytes.decode("utf-8", errors="ignore")
                
        except Exception as e:
            return f"[ফাইল পড়তে ত্রুটি হয়েছে: {str(e)}]"
